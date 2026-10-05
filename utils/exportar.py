import os
from datetime import datetime
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from config import EXCEL_DIR, PDF_DIR
from database.crud import obtener_todos_pacientes


class ReportePDF(FPDF):
    def __init__(self, titulo="Lista de Pacientes"):
        super().__init__('P', 'mm', 'Letter')
        self.titulo = titulo
        self.set_auto_page_break(auto=True, margin=15)

    def header(self):
        self.set_font("Arial", 'B', 16)
        self.set_text_color(25, 118, 210)
        self.cell(0, 10, self.titulo, 0, 1, 'C')
        self.set_draw_color(25, 118, 210)
        self.line(10, 18, 200, 18)
        self.set_font("Arial", 'I', 9)
        self.set_text_color(100)
        self.cell(0, 5, f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}", 0, 1, 'C')
        self.ln(8)

    def footer(self):
        self.set_y(-15)
        self.set_font("Arial", 'I', 8)
        self.set_text_color(150)
        self.cell(0, 10, f'Página {self.page_no()}', 0, 0, 'C')


def _v(c, clave, default=""):
    try:
        valor = c[clave]
        return valor if valor is not None else default
    except (KeyError, IndexError, TypeError):
        return default


def exportar_excel(filtro_sala=None):
    try:
        pacientes = obtener_todos_pacientes()
        if not pacientes:
            return False, "No hay pacientes para exportar"

        datos = []
        for p in pacientes:
            sala = _v(p, 'sala', 0)
            if filtro_sala is not None and filtro_sala != "Todas":
                if filtro_sala == "Sin Sala" and sala != 0:
                    continue
                elif filtro_sala != "Sin Sala" and sala != filtro_sala:
                    continue
            datos.append(p)

        if not datos:
            return False, "No hay pacientes con el filtro seleccionado"

        wb = Workbook()
        ws = wb.active
        ws.title = "Pacientes"

        headers = ['Documento', 'Nombre', 'F. Nac.', 'Género', 'Teléfono',
                   'Dirección', 'Tipo Sangre', 'Alergias', 'Antecedentes',
                   'Presión', 'Temp.', 'Peso', 'Altura', 'Diagnóstico',
                   'Tratamiento', 'Sala', 'Estado', 'F. Registro']

        header_font = Font(bold=True, color="FFFFFF", size=10)
        header_fill = PatternFill(start_color="1976D2", end_color="1976D2", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        cell_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        for row_num, p in enumerate(datos, 2):
            ws.cell(row=row_num, column=1, value=_v(p, 'documento'))
            ws.cell(row=row_num, column=2, value=_v(p, 'nombre'))
            ws.cell(row=row_num, column=3, value=_v(p, 'fecha_nacimiento'))
            ws.cell(row=row_num, column=4, value=_v(p, 'genero'))
            ws.cell(row=row_num, column=5, value=_v(p, 'telefono'))
            ws.cell(row=row_num, column=6, value=_v(p, 'direccion'))
            ws.cell(row=row_num, column=7, value=_v(p, 'tipo_sangre'))
            ws.cell(row=row_num, column=8, value=_v(p, 'alergias'))
            ws.cell(row=row_num, column=9, value=_v(p, 'antecedentes'))
            ws.cell(row=row_num, column=10, value=_v(p, 'presion_arterial'))
            ws.cell(row=row_num, column=11, value=_v(p, 'temperatura'))
            ws.cell(row=row_num, column=12, value=_v(p, 'peso'))
            ws.cell(row=row_num, column=13, value=_v(p, 'altura'))
            ws.cell(row=row_num, column=14, value=_v(p, 'diagnostico'))
            ws.cell(row=row_num, column=15, value=_v(p, 'tratamiento'))
            ws.cell(row=row_num, column=16, value=_v(p, 'sala'))
            ws.cell(row=row_num, column=17, value=_v(p, 'estado'))
            ws.cell(row=row_num, column=18, value=_v(p, 'fecha_registro'))
            for col in range(1, 19):
                ws.cell(row=row_num, column=col).alignment = cell_alignment

        for col in range(1, 19):
            max_len = len(str(headers[col - 1]))
            for row in range(2, len(datos) + 2):
                v = ws.cell(row=row, column=col).value
                if v:
                    max_len = max(max_len, len(str(v)))
            ws.column_dimensions[chr(64 + col)].width = min(max_len + 2, 30)

        os.makedirs(EXCEL_DIR, exist_ok=True)
        nombre = f"Pacientes_{datetime.now().strftime('%d%m%Y_%H%M')}.xlsx"
        ruta = os.path.join(EXCEL_DIR, nombre)
        wb.save(ruta)

        if os.name == 'nt':
            os.startfile(ruta)

        return True, f"Exportado a: {nombre}"
    except Exception as e:
        return False, f"Error: {str(e)}"


def exportar_pdf(filtro_sala=None):
    try:
        pacientes = obtener_todos_pacientes()
        if not pacientes:
            return False, "No hay pacientes para exportar"

        datos = []
        for p in pacientes:
            sala = _v(p, 'sala', 0)
            if filtro_sala is not None and filtro_sala != "Todas":
                if filtro_sala == "Sin Sala" and sala != 0:
                    continue
                elif filtro_sala != "Sin Sala" and sala != filtro_sala:
                    continue
            datos.append(p)

        if not datos:
            return False, "No hay pacientes con el filtro seleccionado"

        pdf = ReportePDF("Lista de Pacientes")
        pdf.add_page()

        col_widths = [28, 50, 30, 25, 20, 20, 20]
        headers = ['Documento', 'Nombre', 'Teléfono', 'Sangre', 'Presión', 'Temp.', 'Estado']

        pdf.set_font("Arial", 'B', 9)
        pdf.set_fill_color(25, 118, 210)
        pdf.set_text_color(255, 255, 255)
        for i, h in enumerate(headers):
            pdf.cell(col_widths[i], 8, h, 1, 0, 'C', True)
        pdf.ln()

        pdf.set_font("Arial", '', 8)
        pdf.set_text_color(0, 0, 0)
        fill = False
        for p in datos:
            fila = [
                _v(p, 'documento', '')[:15],
                _v(p, 'nombre', '')[:28],
                _v(p, 'telefono', '—'),
                _v(p, 'tipo_sangre', '—'),
                _v(p, 'presion_arterial', '—'),
                _v(p, 'temperatura', '—'),
                _v(p, 'estado', '—'),
            ]
            for i, item in enumerate(fila):
                pdf.cell(col_widths[i], 6, str(item), 1, 0, 'L', fill)
            pdf.ln()
            fill = not fill

        os.makedirs(PDF_DIR, exist_ok=True)
        nombre = f"Pacientes_{datetime.now().strftime('%d%m%Y_%H%M')}.pdf"
        ruta = os.path.join(PDF_DIR, nombre)
        pdf.output(ruta)

        if os.name == 'nt':
            os.startfile(ruta)

        return True, f"Exportado a: {nombre}"
    except Exception as e:
        return False, f"Error: {str(e)}"